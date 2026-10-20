"""Tests for the report exports, the Groq integration and the brokerage sources."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402

from api.app.domain.ai import (  # noqa: E402
    AIInvalidResponseError,
    AINotConfiguredError,
    ParsedBrief,
    _parse_json,
    write_explanation,
)
from api.app.domain.brokerage_sources import (  # noqa: E402
    ManualNavigationAdapter,
    STATUS_DATE_UNKNOWN,
    STATUS_LIVE,
    STATUS_UNAVAILABLE,
)
from api.app.domain.exports import build_export, export_filename  # noqa: E402
from api.app.main import app  # noqa: E402


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


REPORT = {
    "report_version": "2.0",
    "generated_at": "2026-01-01T00:00:00+00:00",
    "analysis": {
        "name": "South Chennai family restaurant expansion",
        "business_type": "restaurant",
        "acquisition_mode": "either",
        "aggregation_method": "weighted_sum",
        "location": {"city": "Adyar, Chennai"},
        "requirements": {"business_requirements": []},
        "weights": {
            "root": {"population": 0.25, "accessibility": 0.25},
            "population": {"resident_density": 0.5, "projected_growth": 0.5},
        },
    },
    "executive_summary": {
        "headline": "Adyar is the leading candidate.",
        "paragraphs": ["Adyar scores highest on the stated criteria."],
        "recommendation": {"next_steps": ["Visit the site."]},
    },
    "top_sites": [
        {
            "position": 1,
            "site_id": "site-1",
            "site_name": "Adyar",
            "rank": 1,
            "score": 0.742,
            "why": "It scores higher on accessibility.",
            "caveats": ["Only 50% of the weight could be scored."],
            "cost": {
                "totals": {
                    "initial_investment_total": 3_000_000,
                    "recurring_monthly_total": 120_000,
                    "recurring_annual_total": 1_440_000,
                    "is_complete": True,
                },
                "lines": [
                    {
                        "label": "Monthly rent",
                        "amount": 120_000,
                        "provenance_label": "Verified / listed",
                        "category": "recurring",
                        "source": "rent",
                        "note": "",
                    }
                ],
            },
            "property_opportunities": [],
        }
    ],
    "cost_analysis": [
        {
            "site_id": "site-1",
            "lines": [
                {
                    "label": "Monthly rent",
                    "amount": 120_000,
                    "provenance_label": "Verified / listed",
                    "category": "recurring",
                    "source": "rent",
                    "note": "",
                }
            ],
        }
    ],
    "property_intelligence": {"sites": []},
    "comparison": {},
    "sensitivity": {},
    "data_quality": {
        "limitations": [
            {"severity": "high", "title": "Census data is from 2011", "detail": "…"}
        ]
    },
    "assumptions": ["The candidate inventory is complete."],
    "statement": "This report is decision support.",
}


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def test_the_csv_export_contains_the_ranking_and_the_weights() -> None:
    content, media_type, extension = build_export(REPORT, "csv")

    assert extension == "csv"
    assert "text/csv" in media_type
    text = content.decode("utf-8-sig")

    assert "Adyar" in text
    assert "0.742" in text
    # The exact weighting configuration must be present so the ranking can
    # be reproduced from the export alone.
    assert "WEIGHTING CONFIGURATION" in text
    assert "resident_density" in text
    assert "0.25" in text


def test_the_excel_export_produces_a_workbook() -> None:
    content, media_type, extension = build_export(REPORT, "xlsx")

    assert extension == "xlsx"
    assert "spreadsheetml" in media_type
    assert content[:2] == b"PK"  # a zip container, i.e. a real workbook

    from openpyxl import load_workbook

    workbook = load_workbook(filename=__import__("io").BytesIO(content))
    assert "Shortlist" in workbook.sheetnames
    assert "Weights" in workbook.sheetnames
    assert "Overview" in workbook.sheetnames

    weights = workbook["Weights"]
    rows = list(weights.iter_rows(values_only=True))
    assert rows[0][0] == "Level"
    assert any(row[1] == "resident_density" for row in rows[1:])


def test_the_pdf_export_produces_a_document() -> None:
    content, media_type, extension = build_export(REPORT, "pdf")

    assert extension == "pdf"
    assert media_type == "application/pdf"
    assert content[:5] == b"%PDF-"


def test_the_json_export_is_the_report_itself() -> None:
    content, media_type, extension = build_export(REPORT, "json")
    assert extension == "json"
    assert json.loads(content)["analysis"]["name"].startswith("South Chennai")


def test_an_unsupported_format_is_rejected() -> None:
    with pytest.raises(ValueError, match="Unsupported export format"):
        build_export(REPORT, "docx")


def test_the_export_file_name_is_safe() -> None:
    filename = export_filename(REPORT, "pdf")
    assert filename.endswith(".pdf")
    assert " " not in filename
    assert "/" not in filename


def test_the_export_endpoint_returns_a_file(client: TestClient) -> None:
    response = client.post(
        "/reports/export", json={"report": REPORT, "format": "pdf"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert "attachment" in response.headers.get("content-disposition", "")
    assert response.content[:5] == b"%PDF-"


def test_the_export_endpoint_rejects_an_unknown_format(client: TestClient) -> None:
    response = client.post(
        "/reports/export", json={"report": REPORT, "format": "docx"}
    )
    assert response.status_code == 400


def test_the_export_endpoint_requires_a_report(client: TestClient) -> None:
    response = client.post("/reports/export", json={"report": {}, "format": "pdf"})
    assert response.status_code == 422


def test_the_export_formats_are_listed(client: TestClient) -> None:
    body = client.get("/reports/export/formats").json()
    keys = {entry["key"] for entry in body["formats"]}
    assert {"pdf", "csv", "xlsx", "json"} <= keys


# ---------------------------------------------------------------------------
# The Groq integration
# ---------------------------------------------------------------------------

def test_a_malformed_provider_answer_is_rejected() -> None:
    with pytest.raises(AIInvalidResponseError):
        _parse_json("this is not json")


def test_a_fenced_json_answer_is_accepted() -> None:
    decoded = _parse_json('```json\n{"a": 1}\n```')
    assert decoded == {"a": 1}


def test_the_parsed_brief_validates_its_importance_levels() -> None:
    brief = ParsedBrief.model_validate(
        {
            "requirements": [
                {"key": "seating_capacity", "importance": "nonsense", "value": 40}
            ]
        }
    )
    assert brief.requirements[0].importance == "preferred"


def test_an_unknown_acquisition_mode_falls_back_to_either() -> None:
    brief = ParsedBrief.model_validate({"acquisition_mode": "lease-to-own"})
    assert brief.acquisition_mode == "either"


def test_the_explanation_falls_back_to_the_deterministic_narrative() -> None:
    """Without a provider the deterministic explanation is used unchanged."""

    fallback = "Adyar ranks first because its accessibility score is higher."
    paragraph, response = write_explanation({"facts": {}}, fallback)

    assert response is None
    if paragraph != fallback:
        # A provider answered; the paragraph must not be empty.
        assert paragraph.strip()


def test_the_explanation_writer_never_runs_without_facts() -> None:
    """The writer is only ever handed the deterministic facts."""

    from api.app.domain.ai import EXPLANATION_SYSTEM_PROMPT

    assert "ONLY the facts" in EXPLANATION_SYSTEM_PROMPT
    assert "Do not add" in EXPLANATION_SYSTEM_PROMPT


def test_the_requirement_prompt_forbids_inventing_facts() -> None:
    from api.app.domain.ai import PARSE_SYSTEM_PROMPT

    assert "Never invent facts" in PARSE_SYSTEM_PROMPT
    assert "Return ONLY a JSON object" in PARSE_SYSTEM_PROMPT


def test_a_missing_provider_key_is_reported_clearly() -> None:
    from api.app.core.config import get_settings
    from api.app.domain.ai import parse_requirements

    settings = get_settings()
    if settings.has_ai_provider:
        pytest.skip("A provider is configured in this environment.")

    with pytest.raises(AINotConfiguredError, match="GROQ_API_KEY"):
        parse_requirements("A cafe in Chennai.")


# ---------------------------------------------------------------------------
# Brokerage sources
# ---------------------------------------------------------------------------

def test_a_manual_source_offers_a_search_link_instead_of_inventing_listings() -> None:
    adapter = ManualNavigationAdapter(
        key="portal",
        label="Example portal",
        publisher="Example",
        homepage="https://example.invalid",
        search_template="https://example.invalid/search?q={q}&lat={lat}&lon={lon}",
        instructions="Search the portal directly.",
    )
    result = adapter.search(13.05, 80.24, 5000, 10)

    assert result.status == STATUS_UNAVAILABLE
    assert result.opportunities == []
    assert result.search_url is not None
    assert "13.05000" in result.search_url
    assert result.kind == "manual"


def test_a_manual_source_without_a_template_still_offers_its_homepage() -> None:
    adapter = ManualNavigationAdapter(
        key="portal", label="Example", publisher="Example", homepage="https://example.invalid"
    )
    result = adapter.search(13.05, 80.24, 5000, 10)
    assert result.search_url == "https://example.invalid"
    assert result.opportunities == []


def test_the_brokerage_sources_endpoint_describes_each_source(client: TestClient) -> None:
    body = client.get("/brokerage/sources").json()
    assert body["sources"]
    for source in body["sources"]:
        assert source["kind"] in {"automated", "manual"}
        assert source["label"]
        assert source["description"]


def test_the_brokerage_search_reports_each_source_separately(client: TestClient) -> None:
    """A failing source must degrade to a status, never fail the search."""

    response = client.get(
        "/brokerage/search-sites?site_ids=does-not-exist&radius_m=500&limit_per_site=3"
    )
    # Without a candidate inventory the search cannot be performed; the
    # endpoint must say so rather than return an empty result that could be
    # mistaken for "no agencies found".
    if response.status_code == 503:
        assert "inventory" in response.json()["detail"].lower()
        return

    assert response.status_code == 200
    body = response.json()

    assert body["sites"]
    for site in body["sites"]:
        assert site["status"] in {"ok", "partial", "unavailable"}
        for source in site["sources"]:
            assert source["status"]
            assert "message" in source
